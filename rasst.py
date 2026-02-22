import docx, time, logging, Levenshtein
from fuzzywuzzy import fuzz


def mes_time(func, *args, **kwargs):
    start = time.perf_counter()
    result = func(*args, **kwargs)
    end = time.perf_counter()
    return result, end - start


def levenshtein(str_1, str_2):
    n, m = len(str_1), len(str_2)
    if n > m:
        str_1, str_2 = str_2, str_1
        n, m = m, n

    current_row = range(n + 1)
    for i in range(1, m + 1):
        previous_row, current_row = current_row, [i] + [0] * n
        for j in range(1, n + 1):
            add, delete, change = previous_row[j] + 1, current_row[j - 1] + 1, previous_row[j - 1]
            if str_1[j - 1] != str_2[i - 1]:
                change += 1
            current_row[j] = min(add, delete, change)

    return current_row[n]


logging.basicConfig(filename='time.log', filemode='w', level=logging.INFO)

doc1 = docx.Document('Документ для сравнения_1.docx')
doc2 = docx.Document('Документ для сравнения_2.docx')

str1 = "\n".join(p.text for p in doc1.paragraphs)
str2 = "\n".join(p.text for p in doc2.paragraphs)

# fuzzy
res, t = mes_time(fuzz.ratio, str1, str2)
print(res)
logging.info(f'fuzz.ratio выполнена за {t}')

# левинштейн в процентах
res, t = mes_time(Levenshtein.ratio, str1, str2)
print(res * 100)
logging.info(f'Levenshtein.ratio выполнена за {t}')

# расстояние левинштейна
res, t = mes_time(Levenshtein.distance, str1, str2)
print(res)
logging.info(f'Levenshtein.distance выполнена за {t}')

res, t = mes_time(levenshtein, str1, str2)
print(res)
logging.info(f'levenshtein собственный алгоритм выполнен за {t}')
